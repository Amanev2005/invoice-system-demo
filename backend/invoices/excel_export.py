from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
from datetime import datetime


def export_client_statement_excel(statement_data, client_name):
    """Export client statement to Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    
    # Header styling
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    ws['A1'] = f"Statement of Account - {client_name}"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:G1')
    
    # Date range
    ws['A2'] = f"Period: {statement_data['date_range']['start']} to {statement_data['date_range']['end']}"
    ws.merge_cells('A2:G2')
    
    # Headers
    headers = ['Invoice #', 'Date', 'Due Date', 'Amount', 'Paid', 'Balance', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    
    # Data
    row = 5
    for invoice in statement_data['invoices']:
        ws.cell(row, 1, invoice['invoice_number']).border = border
        ws.cell(row, 2, invoice['invoice_date']).border = border
        ws.cell(row, 3, invoice.get('due_date', '')).border = border
        ws.cell(row, 4, float(invoice['total_amount'])).border = border
        ws.cell(row, 5, float(invoice.get('amount_paid', 0))).border = border
        ws.cell(row, 6, float(invoice.get('balance_due', 0))).border = border
        ws.cell(row, 7, invoice['status'].upper()).border = border
        row += 1
    
    # Summary
    row += 1
    ws.cell(row, 6, "Total Outstanding:").font = Font(bold=True)
    ws.cell(row, 7, float(statement_data['total_outstanding'])).font = Font(bold=True)
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"statement_{client_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def export_revenue_excel(payments, month_filter=None):
    """Export revenue data to Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    
    # Header styling
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    title = f"Revenue Report - {month_filter if month_filter else 'All Time'}"
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:F1')
    
    # Headers
    headers = ['Payment Date', 'Invoice #', 'Client', 'Amount', 'Method', 'Reference']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    
    # Data
    row = 4
    total = 0
    for payment in payments:
        ws.cell(row, 1, payment.payment_date.strftime('%Y-%m-%d')).border = border
        ws.cell(row, 2, payment.invoice.invoice_number).border = border
        ws.cell(row, 3, payment.invoice.client_name).border = border
        ws.cell(row, 4, float(payment.amount)).border = border
        ws.cell(row, 5, payment.payment_method).border = border
        ws.cell(row, 6, payment.reference_number or '').border = border
        total += float(payment.amount)
        row += 1
    
    # Total
    row += 1
    ws.cell(row, 3, "Total Revenue:").font = Font(bold=True)
    ws.cell(row, 4, total).font = Font(bold=True, size=12)
    
    # Adjust column widths
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col].width = 15
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"revenue_{month_filter or 'all_time'}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def export_due_excel(invoices):
    """Export total due/outstanding invoices to Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Outstanding"
    
    # Header styling
    header_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    ws['A1'] = "Outstanding Invoices Report"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:G1')
    
    # Headers
    headers = ['Invoice #', 'Client', 'Invoice Date', 'Due Date', 'Total Amount', 'Paid', 'Balance Due']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    
    # Data
    row = 4
    total_due = 0
    for invoice in invoices:
        if invoice.balance_due > 0:
            ws.cell(row, 1, invoice.invoice_number).border = border
            ws.cell(row, 2, invoice.client_name).border = border
            ws.cell(row, 3, invoice.invoice_date.strftime('%Y-%m-%d')).border = border
            ws.cell(row, 4, invoice.due_date.strftime('%Y-%m-%d') if invoice.due_date else '').border = border
            ws.cell(row, 5, float(invoice.total_amount)).border = border
            ws.cell(row, 6, float(invoice.amount_paid or 0)).border = border
            ws.cell(row, 7, float(invoice.balance_due)).border = border
            total_due += float(invoice.balance_due)
            row += 1
    
    # Total
    row += 1
    ws.cell(row, 6, "Total Due:").font = Font(bold=True)
    ws.cell(row, 7, total_due).font = Font(bold=True, size=12)
    
    # Adjust column widths
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        ws.column_dimensions[col].width = 15
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"outstanding_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
